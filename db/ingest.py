"""Validate .xlsx call exports without database access or environment loading.

Default headers describe our normalized contract. The partner profile supports
the supplied six-column Call Log worksheet. Numeric durations are seconds;
naive timestamps are preserved as supplied, without conversion or a timezone
claim. A configured source identifier is used ONLY for within-file duplicate
detection, never assigned to calls.id (its provenance is not yet confirmed).
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
import json
import math
from pathlib import Path
import re
from typing import BinaryIO

from openpyxl import load_workbook


DURATION_FIELDS = ("in_queue_sec", "agent_sec", "abandon_sec", "routing_sec")
FIELDS = ("start_ts", "abandoned", *DURATION_FIELDS)
# Routing is unused by the current estimators. Other duration columns must be
# present, but may be blank when not applicable to that call's outcome.
REQUIRED_FIELDS = ("start_ts", "abandoned", "in_queue_sec", "agent_sec", "abandon_sec")
ISO_TIMESTAMP = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?")
PARTNER_SHEET = "Call Log"
PARTNER_COLUMNS = {
    "start_ts": "Call Date/Time",
    "in_queue_sec": "Inqueue Time",
    "routing_sec": "Routing Time",
    "agent_sec": "Agent Time",
    "abandoned": "Abandon",
    "abandon_sec": "Abandon Time",
}


@dataclass(frozen=True)
class ValidationIssue:
    field: str | None
    code: str
    message: str


@dataclass
class ValidatedRow:
    row_number: int
    record: dict
    source_identifier: str | None = None
    errors: list[ValidationIssue] = field(default_factory=list)

    @property
    def status(self) -> str:
        if any(error.code != "duplicate_identifier" for error in self.errors):
            return "invalid"
        return "duplicate" if self.errors else "valid"


@dataclass
class ValidationResult:
    rows: list[ValidatedRow] = field(default_factory=list)
    file_errors: list[ValidationIssue] = field(default_factory=list)
    duplicate_detection_enabled: bool = False

    @property
    def summary(self) -> dict:
        return {
            "total_rows": len(self.rows),
            **{f"{status}_rows": sum(row.status == status for row in self.rows)
               for status in ("valid", "invalid", "duplicate")},
        }

    @property
    def ok(self) -> bool:
        return bool(self.rows) and not self.file_errors and all(row.status == "valid" for row in self.rows)

    @property
    def valid_records(self) -> list[dict]:
        # A structural error makes every record ineligible.
        return [] if self.file_errors else [row.record for row in self.rows if row.status == "valid"]

    def report(self) -> dict:
        """Safe for CLI/dashboard display: never include source cell values."""
        def issue(error: ValidationIssue) -> dict:
            return {"field": error.field, "code": error.code, "message": error.message}

        return {
            "ok": self.ok,
            "summary": self.summary,
            "duplicate_detection_enabled": self.duplicate_detection_enabled,
            "warnings": [] if self.duplicate_detection_enabled else [
                "Source-ID duplicate detection was not performed. Validation does not establish that this file is safe to import."
            ],
            "file_errors": [issue(error) for error in self.file_errors],
            "row_errors": [
                {"row_number": row.row_number, "status": row.status,
                 "errors": [issue(error) for error in row.errors]}
                for row in self.rows if row.errors
            ],
        }


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _timestamp(value, timestamp_format: str | None) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        value = value.strip()
        if timestamp_format:
            parsed = datetime.strptime(value, timestamp_format)
        elif ISO_TIMESTAMP.fullmatch(value):
            parsed = datetime.fromisoformat(value)
        else:
            raise ValueError
    else:
        # Do not guess Excel serial dates or interpret numbers as Unix time.
        raise ValueError
    if parsed.tzinfo is not None:
        raise ValueError
    return parsed


def _duration(value) -> float | None:
    if _blank(value):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise ValueError
    number = float(value)
    if not math.isfinite(number) or number < 0:
        raise ValueError
    return number


def _abandoned(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in ("true", "1", "y", "false", "0", "n"):
            return normalized in ("true", "1", "y")
    raise ValueError


def _identifier(value) -> str:
    if isinstance(value, str) and value.strip() and not value.strip().startswith("="):
        return value.strip()
    # Excel numeric identifiers beyond 15 digits can already have lost precision.
    if (not isinstance(value, bool) and isinstance(value, (int, float))
            and abs(value) < 10**15 and math.isfinite(value) and value == int(value)):
        return str(int(value))
    raise ValueError


def validate_rows(
    headers: list,
    rows,
    *,
    column_map: dict[str, str] | None = None,
    identifier_column: str | None = None,
    timestamp_format: str | None = None,
    max_rows: int | None = None,
) -> ValidationResult:
    """Validate rows following a header on Excel row 1; retain every data row.

    Mapping direction: normalized database field -> exact source header.
    Text dates default to ISO YYYY-MM-DD HH:MM[:SS]. Other formats must be
    explicit. Native Excel datetime cells are accepted without format guessing.
    All occurrences of a repeated ID are flagged; no first-wins selection.
    Rows with other errors count as invalid even if their ID also repeats.
    """
    result = ValidationResult(duplicate_detection_enabled=identifier_column is not None)

    def file_error(code, message, name=None):
        result.file_errors.append(ValidationIssue(name, code, message))

    if max_rows is not None:
        original_rows = rows

        def bounded_rows():
            for index, values in enumerate(original_rows):
                if index >= max_rows:
                    file_error("too_many_rows", "The workbook exceeds the supported row limit; no records were accepted.")
                    break
                yield values

        rows = bounded_rows()

    def structural_failure():
        # Keep counts and row references meaningful even when headers/config
        # prevent parsing. No record from such a file is considered valid.
        result.rows = [ValidatedRow(number, {}, errors=[ValidationIssue(
            None, "invalid_structure", "Resolve file-level errors before validating this row."
        )]) for number, _ in enumerate(rows, start=2)]
        return result

    mapping = {name: name for name in FIELDS}
    if column_map is not None:
        if (not isinstance(column_map, dict) or any(
            name not in FIELDS or not isinstance(header, str) or not header.strip()
            for name, header in column_map.items()
        )):
            file_error("invalid_mapping", "Map normalized call fields to nonempty Excel header names.")
            return structural_failure()
        mapping.update(column_map)
    if len(set(mapping.values())) != len(mapping):
        file_error("invalid_mapping", "Each normalized field must map to a distinct header.")
        return structural_failure()
    if timestamp_format is not None and not any(token in timestamp_format for token in ("%H", "%I")):
        file_error("invalid_date_format", "The timestamp format must include an hour (%H or %I).")
        return structural_failure()
    # Header names and cell contents are deliberately excluded from error text.
    named_headers = [value for value in headers if not _blank(value)]
    if any(not isinstance(value, str) for value in named_headers) or len(set(named_headers)) != len(named_headers):
        file_error("invalid_headers", "Headers must be unique text labels on the first row.")
        return structural_failure()
    positions = {value: index for index, value in enumerate(headers) if not _blank(value)}
    for name in REQUIRED_FIELDS:
        if mapping[name] not in positions:
            file_error("missing_column", "A required call column is missing; check the mapping.", name)
    if identifier_column is not None and identifier_column not in positions:
        file_error("missing_column", "The configured identifier column is missing.", "identifier")
    if result.file_errors:
        return structural_failure()

    identifier_rows = defaultdict(list)
    for row_number, values in enumerate(rows, start=2):
        row = ValidatedRow(row_number, {})
        result.rows.append(row)

        def error(name, code, message):
            row.errors.append(ValidationIssue(name, code, message))

        def cell(header):
            index = positions.get(header)
            return values[index] if index is not None and index < len(values) else None

        timestamp = cell(mapping["start_ts"])
        if _blank(timestamp):
            error("start_ts", "required", "A call timestamp is required.")
        else:
            try:
                row.record["start_ts"] = _timestamp(timestamp, timestamp_format)
            except (ValueError, TypeError, OverflowError):
                error("start_ts", "invalid_timestamp", "Use an Excel datetime or the configured text format, without a timezone offset.")
        try:
            row.record["abandoned"] = _abandoned(cell(mapping["abandoned"]))
        except ValueError:
            error("abandoned", "invalid_boolean", "Use Y/N, true/false, or 1/0 for abandonment.")
        for name in DURATION_FIELDS:
            try:
                row.record[name] = _duration(cell(mapping[name]))
            except (ValueError, TypeError, OverflowError):
                error(name, "invalid_duration", "Duration must be finite, nonnegative numeric seconds.")

        # Explicit ingestion policy for downstream estimates, stricter than the
        # database's nullable durations. Zero agent time remains allowed.
        if "abandoned" in row.record:
            required = ("abandon_sec",) if row.record["abandoned"] else ("in_queue_sec", "agent_sec")
            for name in required:
                if row.record.get(name) is None and not any(e.field == name for e in row.errors):
                    error(name, "required_for_outcome", "This duration is required for the call's abandonment outcome.")
        if identifier_column is not None:
            try:
                row.source_identifier = _identifier(cell(identifier_column))
                identifier_rows[row.source_identifier].append(row)
            except ValueError:
                error("identifier", "invalid_identifier", "A nonempty text or exact numeric identifier is required; long IDs must be text.")

    if not result.rows:
        file_error("empty_file", "The worksheet contains no call records.")
    for matches in identifier_rows.values():
        if len(matches) > 1:
            for row in matches:
                row.errors.append(ValidationIssue("identifier", "duplicate_identifier", "This identifier occurs more than once in the file; review all occurrences."))
    return result


def parse_excel(source: str | Path | BinaryIO, *, sheet: str | None = None, **validation_options) -> ValidationResult:
    """Read a local .xlsx path or binary upload stream; perform no writes.

    Only the first sheet is read unless a sheet name is provided. Formula cells
    are not evaluated. Exceptions from workbook parsing never expose cell data.
    """
    if isinstance(source, (str, Path)) and Path(source).suffix.lower() != ".xlsx":
        return ValidationResult(file_errors=[ValidationIssue(None, "unsupported_file", "Provide an .xlsx workbook.")])
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception:
        return ValidationResult(file_errors=[ValidationIssue(None, "unreadable_file", "The workbook could not be read; check that it is a valid, accessible .xlsx file.")])
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            return ValidationResult(file_errors=[ValidationIssue(None, "missing_sheet", "The selected worksheet does not exist.")])
        worksheet = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, ())
        if not headers or all(_blank(value) for value in headers):
            return ValidationResult(file_errors=[ValidationIssue(None, "empty_file", "The worksheet needs a header on its first row and call records.")])
        return validate_rows(list(headers), rows, **validation_options)
    except Exception:
        # Some corrupt workbook errors arise lazily while iterating cells.
        return ValidationResult(file_errors=[ValidationIssue(None, "unreadable_file", "The worksheet could not be read completely; no records were accepted.")])
    finally:
        workbook.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate .xlsx call records locally; never connects to a database.")
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--dry-run", action="store_true", required=True, help="Validation only; required to make intent explicit.")
    parser.add_argument("--profile", choices=("normalized", "partner"), default="normalized",
                        help="partner: supplied six-column Call Log worksheet; normalized: generic headers/options.")
    parser.add_argument("--month", help="Optional YYYY-MM month check for the partner profile. Does not certify completeness.")
    parser.add_argument("--sheet", help="Worksheet name (default: first sheet).")
    parser.add_argument("--column-map", type=Path, help="JSON object mapping normalized fields to Excel headers.")
    parser.add_argument("--identifier-column", help="Reliable source ID header; enables within-file duplicate detection.")
    parser.add_argument("--timestamp-format", help="Explicit strptime format for non-ISO text timestamps.")
    args = parser.parse_args(argv)
    if args.profile == "partner":
        if any((args.sheet, args.column_map, args.identifier_column, args.timestamp_format)):
            parser.error("The partner profile supplies the sheet, headers, and no source identifier. Use normalized for custom options.")
        result = parse_partner_excel(args.file, month=args.month)
        print(json.dumps(result.report(), indent=2))
        return 0 if result.ok else 1
    if args.month is not None:
        parser.error("--month is available with --profile partner.")
    mapping = None
    if args.column_map:
        try:
            mapping = json.loads(args.column_map.read_text())
            if not isinstance(mapping, dict):
                raise ValueError
        except (OSError, ValueError):
            parser.error("Column mapping must be a readable JSON object.")
    result = parse_excel(args.file, sheet=args.sheet, column_map=mapping,
                         identifier_column=args.identifier_column, timestamp_format=args.timestamp_format)
    print(json.dumps(result.report(), indent=2))
    return 0 if result.ok else 1


def parse_partner_excel(source: str | Path | BinaryIO, *, month: str | None = None,
                        max_rows: int = 100_000) -> ValidationResult:
    """Validate the observed partner layout without IDs, credentials or writes.

    All five years can be checked without month. Monthly uploads can specify
    YYYY-MM to flag out-of-month records. Passing checks cannot prove the
    original export included all calls or identify duplicate calls without IDs.
    """
    selected = None
    if month is not None:
        try:
            if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
                raise ValueError
            selected = datetime.strptime(month, "%Y-%m")
        except ValueError:
            return ValidationResult(file_errors=[ValidationIssue(None, "invalid_month", "Select a reporting month using YYYY-MM.")])
    result = parse_excel(source, sheet=PARTNER_SHEET, column_map=PARTNER_COLUMNS,
                         identifier_column=None, max_rows=max_rows)
    if selected is not None:
        for row in result.rows:
            timestamp = row.record.get("start_ts")
            if timestamp is not None and (timestamp.year, timestamp.month) != (selected.year, selected.month):
                row.errors.append(ValidationIssue("start_ts", "outside_month", "This call is outside the selected reporting month."))
    return result


if __name__ == "__main__":
    raise SystemExit(main())
